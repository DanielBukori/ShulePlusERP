import secrets
from functools import wraps

from django.contrib import messages
from django.contrib.auth.decorators import login_required
from django.db import transaction
from django.http import HttpResponseForbidden, HttpResponse
from django.shortcuts import get_object_or_404, redirect, render

from .forms import ResultFilterForm, SubjectScoreFilterForm, AcademicStatisticsFilterForm, ReportPublishForm, StudentRegistrationForm
from shule_core.models import AcademicResult, CustomUser, SchoolClass, Student, Subject, Combination
from shule_core.sms_service import send_sms_to_user
from notifications.utils import notify
from audit.utils import log_action


# ============================================================
# RUHUSA: Views hizi za 'academic' app zinaruhusu HEADMASTER,
# SECOND_MASTER, na ACADEMIC pekee. Dashboard za PARENT, STUDENT,
# TEACHER na DISCIPLINE zitakuwa kwenye app tofauti (accounts).
# ============================================================
ACADEMIC_ROLES = ['HEADMASTER', 'SECOND_MASTER', 'ACADEMIC']


def academic_role_required(view_func):
    @wraps(view_func)
    @login_required
    def wrapper(request, *args, **kwargs):
        if request.user.role not in ACADEMIC_ROLES:
            return HttpResponseForbidden("Hauruhusiwi kufikia ukurasa huu.")
        return view_func(request, *args, **kwargs)
    return wrapper


def _get_school(request):
    return getattr(request, 'school', None)


# ============================================================
# NECTA-style Division points. Hizi ni namba zinazotumika zaidi
# kwa kawaida - THIBITISHA na waraka rasmi wa shule/NECTA kabla
# ya kuzitegemea 100% kwenye ripoti rasmi.
# ============================================================
O_LEVEL_POINTS = {'A': 1, 'B': 2, 'C': 3, 'D': 4, 'F': 5}
O_LEVEL_DIVISIONS = [
    (7, 17, 'I'), (18, 21, 'II'), (22, 25, 'III'), (26, 33, 'IV'), (34, 35, '0'),
]
A_LEVEL_POINTS = {'A': 1, 'B': 2, 'C': 3, 'D': 4, 'E': 5, 'S': 6, 'F': 7}
A_LEVEL_DIVISIONS = [
    (3, 9, 'I'), (10, 12, 'II'), (13, 17, 'III'), (18, 19, 'IV'), (20, 21, '0'),
]


def _score_to_grade(score, level):
    if level == 'A':
        if score >= 80: return 'A'
        elif score >= 70: return 'B'
        elif score >= 60: return 'C'
        elif score >= 50: return 'D'
        elif score >= 40: return 'E'
        elif score >= 35: return 'S'
        else: return 'F'
    else:
        if score >= 75: return 'A'
        elif score >= 65: return 'B'
        elif score >= 45: return 'C'
        elif score >= 30: return 'D'
        else: return 'F'


def _calculate_division(grades, level):
    """
    grades: orodha ya herufi za grade za masomo yote aliyofanya mwanafunzi.
    Inachukua 'best 7' (O-Level) au 'best 3' (A-Level) kwa pointi ndogo zaidi,
    kisha kurudisha (jumla_pointi, division).

    KUMBUKA: kwa A-Level, kwa usahihi kamili, ni masomo ya 'principal' pekee
    yanayopaswa kuhesabiwa (siyo General Studies/subsidiary). Subject model
    ya sasa haina field ya kutofautisha hilo - ongeza is_principal = models.
    BooleanField(default=True) kwenye Subject ukitaka usahihi wa NECTA kamili.
    """
    points_map = O_LEVEL_POINTS if level == 'O' else A_LEVEL_POINTS
    divisions = O_LEVEL_DIVISIONS if level == 'O' else A_LEVEL_DIVISIONS
    n_best = 7 if level == 'O' else 3

    fail_points = points_map['F']
    points = sorted(points_map.get(g, fail_points) for g in grades)[:n_best]
    jumla = sum(points)

    for low, high, div in divisions:
        if low <= jumla <= high:
            return jumla, div
    return jumla, '0'


@academic_role_required
def academic_dashboard_view(request):
    school = _get_school(request)
    if not school:
        return HttpResponseForbidden("Hauruhusiwi kufikia ukurasa huu.")

    from django.db.models import Avg, Count

    classes = SchoolClass.objects.filter(school=school).order_by('level', 'name')

    subject_stats = []
    division_counts = {}
    class_averages = []

    # Orodha ya combos ZOTE zilizowahi kuchapishwa (kwa dropdown ya "Ripoti za
    # Nyuma") - mpya juu, zenye rekodi nyingi zaidi kwanza kwa mwaka mmoja.
    available_combos = list(
        AcademicResult.objects.filter(school=school, is_published=True)
        .values('academic_year', 'term', 'exam_type')
        .annotate(n=Count('id'))
        .order_by('-academic_year', '-n')
    )

    # Kama mtumiaji ameteua combo maalum (kupitia dropdown la "Ripoti za
    # Nyuma", thamani ikiwa "mwaka|muhula|aina" moja), itumike hiyo.
    # Vinginevyo, default ni ile ya kwanza (yenye rekodi nyingi zaidi kwa
    # mwaka wa hivi karibuni - "ya sasa").
    selected_combo = request.GET.get('combo', '')

    latest_combo = None
    if selected_combo:
        parts = selected_combo.split('|')
        if len(parts) == 3:
            sel_year, sel_term, sel_exam_type = parts
            for c in available_combos:
                if (str(c['academic_year']) == sel_year and c['term'] == sel_term
                        and c['exam_type'] == sel_exam_type):
                    latest_combo = c
                    break

    if latest_combo is None and available_combos:
        latest_combo = available_combos[0]

    latest_result = latest_combo  # jina la context variable halibadiliki - template inatumia dot-access, inafanya kazi sawa na dict

    if latest_combo:
        year = latest_combo['academic_year']
        term = latest_combo['term']
        exam_type = latest_combo['exam_type']

        subject_stats = list(
            AcademicResult.objects.filter(
                school=school, is_published=True,
                academic_year=year, term=term, exam_type=exam_type,
            )
            .values('subject__name').annotate(avg_score=Avg('score'), total=Count('id'))
            .order_by('-avg_score')
        )

        for c in classes:
            groups = _build_broadsheet(school, c, term, exam_type, year)
            for g in groups:
                if g['rows']:
                    label = f"{c}" + (f" ({g['combination']})" if g['combination'] else "")
                    class_averages.append({'label': label, 'average': g['class_average']})
                    for row in g['rows']:
                        division_counts[row['division']] = division_counts.get(row['division'], 0) + 1

    class_averages.sort(key=lambda x: x['average'], reverse=True)

    context = {
        'school': school,
        'classes': classes,
        'total_classes': classes.count(),
        'total_subjects': Subject.objects.filter(school=school).count(),
        'total_students': Student.objects.filter(school=school).count(),
        'subject_stats': subject_stats,
        'division_counts': division_counts,
        'class_averages': class_averages,
        'latest_result': latest_result,
        'available_combos': available_combos,
    }

    # Takwimu za ziada kwa Headmaster/Second Master - muhtasari wa shule
    # nzima kutoka apps nyingine (discipline, store, accountant, audit).
    if request.user.role in ('HEADMASTER', 'SECOND_MASTER'):
        from discipline.models import DisciplineAlert
        from store.models import StoreStock
        from accountant.models import get_student_debt
        from audit.models import AuditLog

        context['open_discipline_alerts'] = DisciplineAlert.objects.filter(
            school=school, is_resolved=False
        ).count()
        context['low_stock_count'] = sum(
            1 for s in StoreStock.objects.filter(school=school) if s.is_low
        )
        # NOTE: hesabu hii ni ya moja kwa moja (siyo cached) - kwa shule
        # kubwa sana (wanafunzi 1000+), fikiria kuongeza caching baadaye.
        context['total_outstanding_debt'] = sum(
            get_student_debt(s) for s in Student.objects.filter(school=school)
        )
        context['recent_activity'] = AuditLog.objects.filter(school=school).select_related('actor')[:8]
        context['total_teachers'] = CustomUser.objects.filter(
            school=school, role__in=['TEACHER', 'ACADEMIC', 'SECOND_MASTER', 'DISCIPLINE', 'ACCOUNTANT']
        ).count()

    return render(request, 'academic/academic_dashboard.html', context)


def _build_broadsheet(school, school_class, term, exam_type, academic_year):
    """
    Logic kuu ya broadsheet - inatumika na view ya HTML, PDF, na Excel
    zote tatu, ili zisitofautiane kamwe.
    """
    level = school_class.level
    broadsheet_groups = []

    if level == 'A':
        combinations = Combination.objects.filter(
            students__school_class=school_class
        ).distinct()
        groups = [(c, c.name) for c in combinations]
    else:
        groups = [(None, None)]

    for combination, combo_name in groups:
        students_qs = Student.objects.filter(school=school, school_class=school_class)
        if combination:
            students_qs = students_qs.filter(combination=combination)

        if combination:
            subjects = list(Subject.objects.filter(school=school, level=level, combination=combination))
        else:
            subjects = list(Subject.objects.filter(school=school, level=level, combination__isnull=True))

        rows = []
        for student in students_qs:
            results = AcademicResult.objects.filter(
                student=student, subject__in=subjects,
                exam_type=exam_type, term=term, academic_year=academic_year,
            )
            score_by_subject = {r.subject_id: r.score for r in results}

            scores = [score_by_subject[s.id] for s in subjects if s.id in score_by_subject]
            grades = [_score_to_grade(float(sc), level) for sc in scores]

            total = sum(scores) if scores else 0
            average = (total / len(scores)) if scores else 0

            # Division ihesabiwe TU kama masomo YOTE ya darasa/combination
            # hii yana score kwa exam_type hii - la sivyo tunaonyesha '-'
            # (bado haijakamilika), badala ya kuhesabu Division kutoka
            # masomo machache yaliyojazwa hadi sasa (bug ya awali).
            if subjects and len(scores) == len(subjects):
                points, division = _calculate_division(grades, level)
            else:
                points, division = 0, '-'

            rows.append({
                'student': student,
                'scores': [score_by_subject.get(s.id, '-') for s in subjects],
                'total': total,
                'average': round(float(average), 2),
                'points': points,
                'division': division,
            })

        rows.sort(key=lambda r: r['total'], reverse=True)
        for i, row in enumerate(rows, start=1):
            row['position'] = i

        # Position imeshahesabiwa (kwa total) - sasa panga upya MPANGILIO
        # wa kuonyesha jedwali kwa reg_number, siyo kwa position/division/
        # grades. Position ya kila mwanafunzi inabaki sahihi ndani ya row.
        rows.sort(key=lambda r: r['student'].reg_number)

        class_average = round(sum(r['average'] for r in rows) / len(rows), 2) if rows else 0

        broadsheet_groups.append({
            'combination': combo_name,
            'subjects': subjects,
            'rows': rows,
            'class_average': class_average,
            'total_students': len(rows),
        })

    return broadsheet_groups


@academic_role_required
def academic_class_results_view(request, class_id):
    school = _get_school(request)
    if not school:
        return HttpResponseForbidden("Hauruhusiwi kufikia ukurasa huu.")

    school_class = get_object_or_404(SchoolClass, id=class_id, school=school)
    form = ResultFilterForm(request.GET or None, school=school)

    broadsheet_groups = []

    if form.is_valid():
        broadsheet_groups = _build_broadsheet(
            school, school_class,
            form.cleaned_data['term'], form.cleaned_data['exam_type'], form.cleaned_data['academic_year'],
        )

    context = {
        'school': school,
        'school_class': school_class,
        'form': form,
        'broadsheet_groups': broadsheet_groups,
    }
    return render(request, 'academic/class_results.html', context)


@academic_role_required
def academic_subject_analysis_view(request):
    school = _get_school(request)
    if not school:
        return HttpResponseForbidden("Hauruhusiwi kufikia ukurasa huu.")

    form = SubjectScoreFilterForm(request.GET or None, school=school)
    subject_results = []
    stats = {}

    if form.is_valid():
        subject = form.cleaned_data['subject']
        exam_type = form.cleaned_data['exam_type']
        term = form.cleaned_data['term']
        academic_year = form.cleaned_data['academic_year']

        subject_results = AcademicResult.objects.filter(
            subject=subject, exam_type=exam_type, term=term,
            academic_year=academic_year, school=school,
        ).select_related('student').order_by('-score')

        if subject_results:
            scores = [float(r.score) for r in subject_results]
            pass_mark = 30 if subject.level == 'O' else 40
            stats = {
                'highest': max(scores),
                'lowest': min(scores),
                'average': round(sum(scores) / len(scores), 2),
                'total_students': len(scores),
                'passed': sum(1 for s in scores if s >= pass_mark),
            }

    context = {
        'school': school,
        'form': form,
        'subject_results': subject_results,
        'stats': stats,
    }
    return render(request, 'academic/subject_analysis.html', context)


@academic_role_required
def academic_statistics_view(request):
    """
    Takwimu za Jumla: chagua Muhula + Aina ya Mtihani + Mwaka (na Darasa
    la hiari), unapata:
      1. Pie chart - muhtasari wa jumla wa madaraja (A/B/C/D/F...) - shule
         nzima (au darasa moja ukilichagua).
      2. Bar chart ya grouped - kila SOMO kwenye x-axis, na safu (bars)
         za kila daraja ndani ya somo hilo, ili kuonyesha ni somo lipi
         lina wanafunzi wengi wa A/B/C/D/F.

    Daraja la kila result linakokotolewa kwa kutumia _score_to_grade()
    (helper ile ile inayotumika kwenye broadsheet), kutegemea level ya
    somo husika (O-Level ina A/B/C/D/F; A-Level ina A/B/C/D/E/S/F).
    """
    school = _get_school(request)
    if not school:
        return HttpResponseForbidden("Hauruhusiwi kufikia ukurasa huu.")

    form = AcademicStatisticsFilterForm(request.GET or None, school=school)

    GRADE_ORDER = ['A', 'B', 'C', 'D', 'E', 'S', 'F']
    GRADE_COLORS = {
        'A': '#16a34a', 'B': '#2563eb', 'C': '#ca8a04',
        'D': '#ea580c', 'E': '#7c3aed', 'S': '#0891b2', 'F': '#dc2626',
    }

    pie_labels, pie_data, pie_colors = [], [], []
    bar_subjects, bar_datasets = [], []
    total_results = 0

    if form.is_valid():
        term = form.cleaned_data['term']
        exam_type = form.cleaned_data['exam_type']
        academic_year = form.cleaned_data['academic_year']
        school_class = form.cleaned_data.get('school_class')

        results_qs = AcademicResult.objects.filter(
            school=school, term=term, exam_type=exam_type, academic_year=academic_year,
        ).select_related('subject')
        if school_class:
            results_qs = results_qs.filter(student__school_class=school_class)

        grade_totals = {}
        by_subject = {}
        for r in results_qs:
            grade = _score_to_grade(float(r.score), r.subject.level)
            grade_totals[grade] = grade_totals.get(grade, 0) + 1
            by_subject.setdefault(r.subject.name, {})
            by_subject[r.subject.name][grade] = by_subject[r.subject.name].get(grade, 0) + 1
            total_results += 1

        grades_present = [g for g in GRADE_ORDER if g in grade_totals]
        pie_labels = grades_present
        pie_data = [grade_totals[g] for g in grades_present]
        pie_colors = [GRADE_COLORS[g] for g in grades_present]

        bar_subjects = sorted(by_subject.keys())
        for g in grades_present:
            bar_datasets.append({
                'label': g,
                'color': GRADE_COLORS[g],
                'data': [by_subject[s].get(g, 0) for s in bar_subjects],
            })

    context = {
        'school': school,
        'form': form,
        'pie_labels': pie_labels,
        'pie_data': pie_data,
        'pie_colors': pie_colors,
        'bar_subjects': bar_subjects,
        'bar_datasets': bar_datasets,
        'total_results': total_results,
    }
    return render(request, 'academic/statistics.html', context)


@academic_role_required
def missing_marks_report_view(request):
    """
    Ripoti ya wanafunzi ambao BADO hawajajaziwa alama ya somo fulani kwa
    mtihani uliochaguliwa - kwa ajili ya kufuatilia walimu wamalize
    kuingiza CA zao. TOFAUTI na broadsheet (ambayo inahitaji ukamilifu
    kuhesabu Division), hapa lengo ni KUORODHESHA HASA WALIOKOSA - ili
    Academic Master aweze kuwafuatilia walimu husika. Broadsheet na
    Statistics zote mbili tayari zinaendelea kufanya kazi sahihi hata
    kama baadhi ya alama hazijajazwa (huonyesha '-' badala ya kuzuia).
    """
    school = _get_school(request)
    if not school:
        return HttpResponseForbidden("Hauruhusiwi kufikia ukurasa huu.")

    form = AcademicStatisticsFilterForm(request.GET or None, school=school)
    missing_rows = []

    if form.is_valid():
        term = form.cleaned_data['term']
        exam_type = form.cleaned_data['exam_type']
        academic_year = form.cleaned_data['academic_year']
        selected_class = form.cleaned_data.get('school_class')

        classes = [selected_class] if selected_class else SchoolClass.objects.filter(school=school)

        for c in classes:
            level = c.level
            if level == 'A':
                combinations = Combination.objects.filter(students__school_class=c).distinct()
                groups = [(comb, comb.name) for comb in combinations]
            else:
                groups = [(None, None)]

            for combination, combo_name in groups:
                students_qs = Student.objects.filter(school=school, school_class=c)
                if combination:
                    students_qs = students_qs.filter(combination=combination)
                    subjects = list(Subject.objects.filter(school=school, level=level, combination=combination))
                else:
                    subjects = list(Subject.objects.filter(school=school, level=level, combination__isnull=True))

                if not subjects:
                    continue

                for student in students_qs:
                    existing_subject_ids = set(
                        AcademicResult.objects.filter(
                            student=student, subject__in=subjects,
                            exam_type=exam_type, term=term, academic_year=academic_year,
                        ).values_list('subject_id', flat=True)
                    )
                    for subject in subjects:
                        if subject.id not in existing_subject_ids:
                            missing_rows.append({
                                'student': student, 'school_class': c,
                                'combination': combo_name, 'subject': subject,
                            })

    missing_rows.sort(key=lambda r: (str(r['school_class']), r['student'].reg_number, r['subject'].name))
    missing_students_count = len({r['student'].id for r in missing_rows})

    context = {
        'school': school,
        'form': form,
        'missing_rows': missing_rows,
        'missing_count': len(missing_rows),
        'missing_students_count': missing_students_count,
    }
    return render(request, 'academic/missing_marks_report.html', context)


@academic_role_required
def auto_absent_report_view(request):
    """
    Ripoti ya wanafunzi WALIOPEWA F(0) KIOTOMATIKO wakati wa Publish
    (is_auto_absent=True) - TOFAUTI na F halisi aliyoipata mwanafunzi
    kwa kufanya mtihani vibaya. Hii ni kwa ajili ya ufuatiliaji WA
    BAADA YA PUBLISH - mfano: kuhakiki kama mwanafunzi alikuwa
    mgonjwa/hayupo shuleni, au ni kosa la kuandika ambalo bado
    linaweza kurekebishwa kupitia Marks Entry ya kawaida.
    """
    school = _get_school(request)
    if not school:
        return HttpResponseForbidden("Hauruhusiwi kufikia ukurasa huu.")

    form = AcademicStatisticsFilterForm(request.GET or None, school=school)
    auto_absent_results = []

    if form.is_valid():
        term = form.cleaned_data['term']
        exam_type = form.cleaned_data['exam_type']
        academic_year = form.cleaned_data['academic_year']
        selected_class = form.cleaned_data.get('school_class')

        results_qs = AcademicResult.objects.filter(
            school=school, term=term, exam_type=exam_type, academic_year=academic_year,
            is_auto_absent=True,
        ).select_related('student', 'subject', 'student__school_class').order_by(
            'student__school_class__name', 'student__reg_number', 'subject__name'
        )
        if selected_class:
            results_qs = results_qs.filter(student__school_class=selected_class)

        auto_absent_results = list(results_qs)

    auto_absent_students_count = len({r.student_id for r in auto_absent_results})

    context = {
        'school': school,
        'form': form,
        'auto_absent_results': auto_absent_results,
        'auto_absent_count': len(auto_absent_results),
        'auto_absent_students_count': auto_absent_students_count,
    }
    return render(request, 'academic/auto_absent_report.html', context)


@academic_role_required
def academic_publish_results_view(request):
    """
    Kuchapisha matokeo ya darasa. KABLA ya kuchapisha, mfumo unafanya
    ukaguzi wa 'wamekosa kufanya mtihani uliokuwepo':

    - Kwa kila somo, kama LIPO angalau alama MOJA iliyorekodiwa kwa
      darasa/combination hii (yaani mtihani wa somo hilo ULIFANYIKA),
      basi mwanafunzi YEYOTE wa darasa hilo asiye na alama ya somo
      hilo anapewa alama 0 (= daraja F) KIOTOMATIKO - kama alikosa
      mtihani bila sababu ya maandishi.
    - Kama somo halina alama YOYOTE kwa darasa hilo (yaani mtihani wa
      somo hilo HAUKUFANYIKA kabisa mzunguko huu), halifanyiwi lolote -
      halitaonekana popote kwenye ripoti kama kwamba mtihani haukuwepo.
    """
    school = _get_school(request)
    if not school:
        return HttpResponseForbidden("Hauruhusiwi kufikia ukurasa huu.")

    published = False
    auto_f_count = 0

    if request.method == 'POST':
        form = ReportPublishForm(request.POST, school=school)
        if form.is_valid():
            school_class = form.cleaned_data['school_class']
            term = form.cleaned_data['term']
            exam_type = form.cleaned_data['exam_type']
            academic_year = form.cleaned_data['academic_year']
            level = school_class.level

            with transaction.atomic():
                if level == 'A':
                    combinations = Combination.objects.filter(
                        students__school_class=school_class
                    ).distinct()
                    groups = [(c, c.name) for c in combinations]
                else:
                    groups = [(None, None)]

                for combination, combo_name in groups:
                    students_qs = Student.objects.filter(school=school, school_class=school_class)
                    if combination:
                        students_qs = students_qs.filter(combination=combination)
                        subjects = list(Subject.objects.filter(school=school, level=level, combination=combination))
                    else:
                        subjects = list(Subject.objects.filter(school=school, level=level, combination__isnull=True))

                    for subject in subjects:
                        base_filter = dict(
                            school=school, subject=subject, school_class=school_class,
                            term=term, exam_type=exam_type, academic_year=academic_year,
                            combination=combination,
                        )
                        # Je, somo hili LILIFANYIWA mtihani (angalau alama moja
                        # tayari ipo)? Kama HAPANA - ruka kabisa somo hili;
                        # mtihani haukufanyika, si 'kila mtu amekosa'.
                        if not AcademicResult.objects.filter(**base_filter).exists():
                            continue

                        existing_student_ids = set(
                            AcademicResult.objects.filter(**base_filter).values_list('student_id', flat=True)
                        )
                        for student in students_qs:
                            if student.id not in existing_student_ids:
                                AcademicResult.objects.create(
                                    school=school, student=student, subject=subject,
                                    combination=combination, school_class=school_class,
                                    exam_type=exam_type, term=term, academic_year=academic_year,
                                    score=0, recorded_by=request.user, is_auto_absent=True,
                                )
                                auto_f_count += 1

                AcademicResult.objects.filter(
                    school=school, student__school_class=school_class,
                    term=term, exam_type=exam_type, academic_year=academic_year,
                ).update(is_published=True)

            published = True

            # Tuma SMS kwa wazazi wote wa darasa hili wenye akaunti
            students = Student.objects.filter(school=school, school_class=school_class).select_related('parent')
            message = (
                f"ShulePlusERP: Matokeo ya {school_class} ({exam_type}, {term}) "
                f"yamechapishwa. Ingia mfumoni kuona matokeo ya mtoto wako."
            )
            for student in students:
                if student.parent:
                    notify(student.parent, school, message, link='/family/parent/dashboard/')

            log_msg = f"Amechapisha matokeo ya {school_class} - {exam_type}, {term}, {academic_year}."
            if auto_f_count:
                log_msg += f" Alama {auto_f_count} za 'F' ziliwekwa kiotomatiki kwa waliokosa mtihani uliokuwepo."
            log_action(request.user, school, 'RESULTS_PUBLISHED', log_msg)
    else:
        form = ReportPublishForm(school=school)

    context = {'school': school, 'form': form, 'published': published, 'auto_f_count': auto_f_count}
    return render(request, 'academic/publish_results.html', context)


@academic_role_required
def register_student_view(request):
    """
    Sajili mwanafunzi mpya, na kiotomatiki:
      - tengeneza/unganisha akaunti ya Mzazi (username = namba ya simu)
      - tengeneza akaunti ya Mwanafunzi (username = reg_number, password = reg_number)
    Username zote zina prefix ya subdomain ya shule ili zisigongane
    na shule nyingine (CustomUser.username ni unique kwa mfumo mzima).
    """
    school = _get_school(request)
    if not school:
        return HttpResponseForbidden("Hauruhusiwi kufikia ukurasa huu.")

    if request.method == 'POST':
        form = StudentRegistrationForm(request.POST, school=school)
        if form.is_valid():
            with transaction.atomic():
                student = form.save(commit=False)
                student.school = school

                phone_digits = ''.join(ch for ch in form.cleaned_data['parent_phone'] if ch.isdigit())
                parent_username = f"{school.subdomain}{phone_digits}"

                parent_user = CustomUser.objects.filter(
                    school=school, role='PARENT', phone=phone_digits
                ).first()

                parent_temp_password = None
                if not parent_user:
                    parent_temp_password = f"{secrets.randbelow(900000) + 100000}"  # PIN ya tarakimu 6
                    parent_user = CustomUser.objects.create_user(
                        username=parent_username,
                        password=parent_temp_password,
                        email=f"{parent_username}@no-email.shulepluserp.local",
                        phone=phone_digits,
                        role='PARENT',
                        school=school,
                        first_name=form.cleaned_data['parent_name'],
                        must_change_password=True,
                    )
                student.parent = parent_user

                student_username = f"{school.subdomain}{student.reg_number}"
                student_user = CustomUser.objects.create_user(
                    username=student_username,
                    password=student.reg_number,
                    email=f"{student_username}@no-email.shulepluserp.local",
                    role='STUDENT',
                    school=school,
                    first_name=student.first_name,
                    last_name=student.last_name,
                    must_change_password=True,
                )
                student.user = student_user
                student.save()

            success_msg = (
                f"Mwanafunzi {student} amesajiliwa kikamilifu. "
                f"Akaunti ya mwanafunzi: {student_username} / {student.reg_number}. "
            )
            if parent_temp_password:
                success_msg += (
                    f"Akaunti mpya ya mzazi: {parent_username} / {parent_temp_password} "
                    f"(mtumie kwa SMS sasa hivi)."
                )
            else:
                success_msg += "Mzazi tayari ana akaunti (ameshaunganishwa na mtoto mwingine)."

            messages.success(request, success_msg)
            log_action(
                request.user, school, 'STUDENT_REGISTERED',
                f"Amesajili mwanafunzi {student} ({student.reg_number}) - {student.school_class}."
            )
            return redirect('academic:register_student')
    else:
        form = StudentRegistrationForm(school=school)

    context = {'school': school, 'form': form}
    return render(request, 'academic/register_student.html', context)


@academic_role_required
def download_broadsheet_pdf(request, class_id):
    """
    PDF ya broadsheet - inatumia template TOFAUTI (class_results_pdf.html,
    CSS halisi siyo Tailwind CDN, kwa sababu WeasyPrint haitumii JS).
    """
    import weasyprint

    school = _get_school(request)
    if not school:
        return HttpResponseForbidden("Hauruhusiwi kufikia ukurasa huu.")

    school_class = get_object_or_404(SchoolClass, id=class_id, school=school)
    form = ResultFilterForm(request.GET or None, school=school)

    if not form.is_valid():
        return HttpResponseForbidden("Chagua Term/Exam/Academic Year kwanza kabla ya ku-download.")

    broadsheet_groups = _build_broadsheet(
        school, school_class,
        form.cleaned_data['term'], form.cleaned_data['exam_type'], form.cleaned_data['academic_year'],
    )

    html_string = render(request, 'academic/class_results_pdf.html', {
        'school': school, 'school_class': school_class, 'broadsheet_groups': broadsheet_groups,
        'term': form.cleaned_data['term'], 'exam_type': form.cleaned_data['exam_type'],
        'academic_year': form.cleaned_data['academic_year'],
    }).content.decode('utf-8')

    pdf_file = weasyprint.HTML(string=html_string, base_url=request.build_absolute_uri('/')).write_pdf()

    response = HttpResponse(pdf_file, content_type='application/pdf')
    filename = f"matokeo_{school_class}_{form.cleaned_data['academic_year']}.pdf".replace(' ', '_')
    response['Content-Disposition'] = f'attachment; filename="{filename}"'
    return response


@academic_role_required
def download_broadsheet_excel(request, class_id):
    """Excel (.xlsx) ya broadsheet - openpyxl, safu moja kwa kila combination."""
    import openpyxl
    from openpyxl.styles import Font, PatternFill, Alignment

    school = _get_school(request)
    if not school:
        return HttpResponseForbidden("Hauruhusiwi kufikia ukurasa huu.")

    school_class = get_object_or_404(SchoolClass, id=class_id, school=school)
    form = ResultFilterForm(request.GET or None, school=school)

    if not form.is_valid():
        return HttpResponseForbidden("Chagua Term/Exam/Academic Year kwanza kabla ya ku-download.")

    broadsheet_groups = _build_broadsheet(
        school, school_class,
        form.cleaned_data['term'], form.cleaned_data['exam_type'], form.cleaned_data['academic_year'],
    )

    wb = openpyxl.Workbook()
    wb.remove(wb.active)

    header_fill = PatternFill(start_color='0B2545', end_color='0B2545', fill_type='solid')
    header_font = Font(color='FFFFFF', bold=True, size=10)

    for group in broadsheet_groups:
        sheet_name = (group['combination'] or str(school_class))[:31]
        ws = wb.create_sheet(title=sheet_name)

        headers = ['Nafasi', 'Reg Number', 'Jina'] + [s.code for s in group['subjects']] + \
                  ['Jumla', 'Wastani', 'Pointi', 'Division']
        ws.append(headers)
        for cell in ws[1]:
            cell.fill = header_fill
            cell.font = header_font
            cell.alignment = Alignment(horizontal='center')

        for row in group['rows']:
            ws.append([
                row['position'], row['student'].reg_number,
                f"{row['student'].first_name} {row['student'].last_name}",
                *row['scores'], row['total'], row['average'], row['points'], row['division'],
            ])

        for col in ws.columns:
            max_len = max((len(str(c.value)) for c in col if c.value is not None), default=8)
            ws.column_dimensions[col[0].column_letter].width = max_len + 3

    response = HttpResponse(content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')
    filename = f"matokeo_{school_class}_{form.cleaned_data['academic_year']}.xlsx".replace(' ', '_')
    response['Content-Disposition'] = f'attachment; filename="{filename}"'
    wb.save(response)
    return response


@academic_role_required
def download_import_template_view(request):
    """Excel template tupu (headers + mfano mmoja) kwa Bulk Student Import."""
    import openpyxl
    from openpyxl.styles import Font, PatternFill

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Wanafunzi"

    headers = [
        'first_name', 'middle_name', 'last_name', 'reg_number',
        'school_class', 'combination', 'parent_name', 'parent_phone',
    ]
    ws.append(headers)
    header_fill = PatternFill(start_color='0B2545', end_color='0B2545', fill_type='solid')
    for cell in ws[1]:
        cell.fill = header_fill
        cell.font = Font(color='FFFFFF', bold=True)

    # Mfano - mstari huu ni maelekezo tu, futa kabla ya kupakia
    ws.append([
        'Juma', 'Ally', 'Mussa', 'MWG/2026/001',
        'Form I', '', 'Ally Mussa', '0712345678',
    ])
    ws.append([
        'Neema', '', 'Josephat', 'MWG/2026/002',
        'Form V', 'PCM', 'Josephat Method', '0755000111',
    ])

    for col in ws.columns:
        ws.column_dimensions[col[0].column_letter].width = 18

    response = HttpResponse(content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')
    response['Content-Disposition'] = 'attachment; filename="template_wanafunzi.xlsx"'
    wb.save(response)
    return response


@academic_role_required
def bulk_register_students_view(request):
    """
    Sajili wanafunzi WENGI mara moja kutoka Excel - kila mstari unatengeneza
    Student + akaunti ya Mwanafunzi + akaunti/kuunganisha Mzazi, kama
    register_student_view moja moja - ila kwa mkupuo.
    """
    import openpyxl

    school = _get_school(request)
    if not school:
        return HttpResponseForbidden("Hauruhusiwi kufikia ukurasa huu.")

    created = []
    errors = []

    if request.method == 'POST' and request.FILES.get('excel_file'):
        excel_file = request.FILES['excel_file']
        try:
            wb = openpyxl.load_workbook(excel_file)
            ws = wb.active
        except Exception:
            messages.error(request, "Faili hii si Excel sahihi (.xlsx). Jaribu tena.")
            return render(request, 'academic/bulk_register_students.html', {'school': school})

        rows = list(ws.iter_rows(min_row=2, values_only=True))

        for i, row in enumerate(rows, start=2):
            if not row or not any(row):
                continue  # mstari mtupu - ruka

            try:
                (first_name, middle_name, last_name, reg_number,
                 class_name, combo_name, parent_name, parent_phone) = (row + (None,) * 8)[:8]

                if not all([first_name, last_name, reg_number, class_name, parent_name, parent_phone]):
                    errors.append(f"Mstari {i}: kuna column muhimu iliyoachwa wazi.")
                    continue

                reg_number = str(reg_number).strip()
                if Student.objects.filter(school=school, reg_number=reg_number).exists():
                    errors.append(f"Mstari {i}: Reg Number '{reg_number}' tayari ipo.")
                    continue

                school_class = SchoolClass.objects.filter(school=school, name__iexact=str(class_name).strip()).first()
                if not school_class:
                    errors.append(f"Mstari {i}: Darasa '{class_name}' halipo kwenye mfumo.")
                    continue

                combination = None
                if school_class.level == 'A':
                    if not combo_name:
                        errors.append(f"Mstari {i}: Darasa la A-Level linahitaji 'combination'.")
                        continue
                    combination = Combination.objects.filter(name__iexact=str(combo_name).strip()).first()
                    if not combination:
                        errors.append(f"Mstari {i}: Combination '{combo_name}' haipo.")
                        continue

                phone_digits = ''.join(ch for ch in str(parent_phone) if ch.isdigit())
                if len(phone_digits) < 9:
                    errors.append(f"Mstari {i}: Namba ya simu ya mzazi si sahihi.")
                    continue

                with transaction.atomic():
                    parent_username = f"{school.subdomain}{phone_digits}"
                    parent_user = CustomUser.objects.filter(school=school, role='PARENT', phone=phone_digits).first()
                    parent_temp_password = None
                    if not parent_user:
                        parent_temp_password = f"{secrets.randbelow(900000) + 100000}"
                        parent_user = CustomUser.objects.create_user(
                            username=parent_username, password=parent_temp_password,
                            email=f"{parent_username}@no-email.shulepluserp.local",
                            phone=phone_digits, role='PARENT', school=school,
                            first_name=str(parent_name), must_change_password=True,
                        )

                    student_username = f"{school.subdomain}{reg_number}"
                    student_user = CustomUser.objects.create_user(
                        username=student_username, password=reg_number,
                        email=f"{student_username}@no-email.shulepluserp.local",
                        role='STUDENT', school=school,
                        first_name=str(first_name), last_name=str(last_name),
                        must_change_password=True,
                    )

                    student = Student.objects.create(
                        school=school, user=student_user, school_class=school_class,
                        combination=combination, first_name=str(first_name),
                        middle_name=str(middle_name) if middle_name else None,
                        last_name=str(last_name), reg_number=reg_number,
                        parent=parent_user, parent_name=str(parent_name),
                        parent_phone=phone_digits,
                    )

                created.append({
                    'student': student, 'student_username': student_username,
                    'parent_username': parent_username if parent_temp_password else None,
                    'parent_password': parent_temp_password,
                })
            except Exception as e:
                errors.append(f"Mstari {i}: hitilafu isiyotarajiwa - {e}")

        if created:
            log_action(
                request.user, school, 'STUDENT_REGISTERED',
                f"Bulk import: amesajili wanafunzi {len(created)} kwa mkupuo (Excel)."
            )
            messages.success(request, f"Wanafunzi {len(created)} wamesajiliwa kikamilifu.")
        if errors:
            messages.warning(request, f"Mistari {len(errors)} imeruka - angalia maelezo hapa chini.")

    context = {'school': school, 'created': created, 'errors': errors}
    return render(request, 'academic/bulk_register_students.html', context)


@academic_role_required
def promote_students_view(request):
    """
    Kupandisha wanafunzi darasa mwishoni mwa mwaka - Academic/Headmaster
    anachagua kila darasa 'linapanda kwenda wapi' (au 'Wamehitimu' kama
    ni la mwisho, mfano Form IV kabla ya A-Level - hilo linahitaji
    usajili mpya wa combination kwa hiyo halifanyiki kiotomatiki hapa).
    """
    school = _get_school(request)
    if not school:
        return HttpResponseForbidden("Hauruhusiwi kufikia ukurasa huu.")

    from django.db.models import Count
    classes = SchoolClass.objects.filter(school=school).order_by('level', 'name').annotate(
        student_count=Count('students')
    )

    if request.method == 'POST':
        promoted_summary = []
        with transaction.atomic():
            for c in classes:
                choice = request.POST.get(f'promote_{c.id}', '').strip()
                if not choice or choice == 'NONE':
                    continue
                students_qs = Student.objects.filter(school_class=c)
                count = students_qs.count()
                if count == 0:
                    continue

                if choice == 'GRADUATE':
                    students_qs.update(school_class=None)
                    promoted_summary.append(f"{c}: wanafunzi {count} wamehitimu (wametolewa kwenye darasa).")
                else:
                    target = SchoolClass.objects.filter(school=school, id=choice).first()
                    if target:
                        students_qs.update(school_class=target)
                        promoted_summary.append(f"{c}: wanafunzi {count} wamepandishwa kwenda {target}.")

        if promoted_summary:
            log_action(
                request.user, school, 'STUDENT_REGISTERED',
                "Promotion: " + " | ".join(promoted_summary)
            )
            messages.success(request, "Promotion imekamilika: " + " | ".join(promoted_summary))
        else:
            messages.warning(request, "Hakuna darasa lililochaguliwa kupandishwa.")
        return redirect('academic:promote_students')

    context = {'school': school, 'classes': classes}
    return render(request, 'academic/promote_students.html', context)


@academic_role_required
def global_search_view(request):
    """Tafuta mwanafunzi (jina/reg_number) au mwalimu (jina) kwa haraka."""
    school = _get_school(request)
    if not school:
        return HttpResponseForbidden("Hauruhusiwi kufikia ukurasa huu.")

    query = request.GET.get('q', '').strip()
    students_found = []
    staff_found = []

    if query:
        from django.db.models import Q
        students_found = Student.objects.filter(school=school).filter(
            Q(first_name__icontains=query) | Q(last_name__icontains=query) | Q(reg_number__icontains=query)
        ).select_related('school_class')[:20]

        staff_found = CustomUser.objects.filter(school=school).filter(
            Q(first_name__icontains=query) | Q(last_name__icontains=query) | Q(username__icontains=query)
        ).exclude(role__in=['PARENT', 'STUDENT'])[:20]

    context = {'school': school, 'query': query, 'students_found': students_found, 'staff_found': staff_found}
    return render(request, 'academic/global_search.html', context)
